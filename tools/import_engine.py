"""Entry point import data master: run_import (dipanggil wizard, sf.import.history,
dan post_init_hook modul klien lewat import_bundled_data), validasi & cleanup
sebelum import. Importer per sheet ada di modul lain di folder tools/ ini.
"""

import glob
import logging
import os
from datetime import timedelta

from openpyxl import load_workbook

from odoo.exceptions import UserError

from .common import MODULE, find_sheet, _parse_sheet_date, ImportLogger, _sheet_rows
from .company import (
    _read_decimal_accuracy_percentage_analytic,
    _apply_analytic_percentage_precision,
    _import_company,
)
from .account import _import_account_account, _import_account_journal
from .partner import (
    _import_res_partner,
    _import_partner_default_accounts,
    _import_partner_opening_balance,
)
from .opening_move import _import_opening_move, _import_kas_bank, _reconcile_liquidity_transfer
from .asset import _import_account_asset_model, _import_account_asset
from .report_analytic import (
    _import_account_analytic_plan,
    _import_account_analytic_account,
    _import_account_report,
    _import_account_report_line,
)

# Diekspor ulang untuk pemanggil lama (hooks.py, models/import_history.py)
# yang meng-import nama-nama ini dari import_engine.
from .common import SHEET_ALIASES  # noqa: F401
from .account import ACCOUNT_TYPES, JOURNAL_TYPES  # noqa: F401
from .asset import ASSET_METHODS, ASSET_PERIODS  # noqa: F401

_logger = logging.getLogger(__name__)


def find_data_file(data_dir):
    """Path to the single .xlsx under data_dir, or None if there isn't one -
    the data import is optional seed data, not a structural part of a client
    module, so a missing file just skips the import instead of blocking
    install. More than one file is treated as a real misconfiguration
    (which one would we import?) and still raises.
    """
    matches = glob.glob(os.path.join(data_dir, "*.xlsx"))
    if len(matches) > 1:
        raise ValueError(f"Expected at most one .xlsx file in {data_dir}, found {len(matches)}")
    return matches[0] if matches else None


def import_bundled_data(env, module_dir):
    """Convenience entry point for a client module's post_init_hook: find
    the single .xlsx under <module_dir>/data and run the import, or skip
    with a warning if there isn't one.
    """
    data_dir = os.path.join(module_dir, "data")
    path = find_data_file(data_dir)
    if not path:
        _logger.warning("No .xlsx master data file found under %s - skipping data import.", data_dir)
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    wb_formulas = load_workbook(path, read_only=True, data_only=False)
    run_import(env, wb, wb_formulas=wb_formulas)


def _read_opening_balance_date(wb):
    """Read OPENING_BALANCE_DATE, a key/value row in the "company" sheet
    (falls back to the older "petunjuk" location for workbooks that haven't
    been migrated yet). Ini tanggal cutover opening balance yang sebenarnya
    (mis. "saldo per 30 Juni 2026") - fiscal_year_start di run_import()
    di-derive dari sini (+1 hari), bukan sebaliknya, karena inilah nilai
    yang paling natural diisi orang: tanggal per kapan saldo awal berlaku.
    """
    for sheet_key in ("company", "c", "petunjuk", "p"):
        actual_sheet = find_sheet(wb, sheet_key)
        if not actual_sheet or actual_sheet not in wb.sheetnames:
            continue
        for row in wb[actual_sheet].iter_rows(values_only=True):
            if row and row[0] == "OPENING_BALANCE_DATE":
                return _parse_sheet_date(row[1], "OPENING_BALANCE_DATE")
    raise ValueError('"OPENING_BALANCE_DATE" tidak ditemukan di sheet "company" (atau "c") atau "petunjuk" (atau "p")')


def _warn_uncomputed_formulas(wb_formulas, wb_values, logger=None):
    """data_only=True (dipakai di run_import) hanya bisa membaca *cached value*
    terakhir dari sel formula, bukan rumusnya - itu cukup untuk file yang
    sudah pernah dihitung ulang oleh Excel/Google Sheets (kasus normal, lihat
    _download_google_sheet yang selalu export dalam kondisi ter-kalkulasi).
    Tapi kalau file .xlsx dibuat/diedit oleh tool lain yang menulis string
    rumus tanpa pernah menghitungnya, cache-nya kosong dan openpyxl
    mengembalikan None - baris/kolom itu akan diam-diam diperlakukan sebagai
    blank oleh importer. Fungsi ini membandingkan workbook rumus
    (data_only=False) dengan workbook nilai (data_only=True) untuk
    mendeteksi kasus itu secara eksplisit dan mencatatnya sebagai warning,
    alih-alih membiarkannya lolos sebagai data kosong tanpa jejak.
    """
    for sheet_name in wb_formulas.sheetnames:
        if sheet_name not in wb_values.sheetnames:
            continue
        ws_formulas = wb_formulas[sheet_name]
        ws_values = wb_values[sheet_name]
        for row_formulas, row_values in zip(ws_formulas.iter_rows(), ws_values.iter_rows()):
            for cell_formula, cell_value in zip(row_formulas, row_values):
                if cell_formula.data_type == "f" and cell_value.value is None:
                    message = (
                        f'Sel {cell_formula.coordinate} pada sheet "{sheet_name}" berisi rumus '
                        f"({cell_formula.value}) yang belum pernah dihitung (cached value kosong) - "
                        "dibaca sebagai kosong oleh importer. Buka & simpan ulang file-nya di "
                        "Excel/Google Sheets supaya nilainya ikut ter-hitung sebelum diimport."
                    )
                    _logger.warning(message)
                    if logger:
                        logger.log(sheet_name, cell_formula.row, cell_formula.coordinate, "", "warning", message)
    if logger:
        logger.flush()


def _check_user_journal_entries(env, company):
    """Cek apakah sudah ada journal entries operasional yang diinput oleh user.
    Mengembalikan recordset account.move buatan user jika ditemukan.
    """
    opening_move_id = company.account_opening_move_id.id if company.account_opening_move_id else None

    domain = [("company_id", "=", company.id)]
    moves = env["account.move"].search(domain)

    user_moves = env["account.move"]
    for m in moves:
        # Lewati opening journal entry neraca bawaan/sistem
        if opening_move_id and m.id == opening_move_id:
            continue

        # Lewati opening bank statement lines dari import saldo awal
        if m.statement_line_id:
            ref = (m.statement_line_id.payment_ref or "").lower()
            if "saldo awal" in ref or "opening" in ref:
                continue

        # Lewati opening vendor_bill / customer_invoice yang dibuat oleh import ini
        imd = env["ir.model.data"].search([
            ("model", "=", "account.move"),
            ("res_id", "=", m.id),
            ("module", "in", (MODULE, "__import__")),
        ], limit=1)
        if imd:
            continue

        user_moves |= m

    return user_moves


def _cleanup_previous_data(env, wb, company):
    """Membersihkan sisa data master dan saldo awal dari import sebelumnya
    atau sisa kloning database ketika belum ada transaksi operasional user.
    """
    _logger.info("Membersihkan data sisa import sebelumnya untuk company %s...", company.name)

    # 0. Bersihkan mapping ir.model.data dinamis lama agar ID baru tidak tertukar/overwrite
    imds = env["ir.model.data"].search([
        ("module", "=", MODULE),
        ("model", "in", ("account.account", "account.journal", "account.asset", "account.bank.statement.line")),
    ])
    if imds:
        imds.unlink()

    # 1. Bersihkan statement lines saldo awal lama
    st_lines = env["account.bank.statement.line"].search([("company_id", "=", company.id)])
    for st in st_lines:
        if st.is_reconciled:
            try:
                st.line_ids.remove_move_reconcile()
            except Exception:
                pass
    if st_lines:
        st_lines.unlink()

    # 2. Reset opening balance move neraca lama
    if company.account_opening_move_id:
        op_move = company.account_opening_move_id
        if op_move.state == "posted":
            op_move.button_draft()
        op_move.line_ids.unlink()

    # 3. Bersihkan jurnal custom lama yang TIDAK ADA di spreadsheet baru
    sheet_journal_codes = set()
    sheet_journal_names = set()
    journal_sheet = find_sheet(wb, "account.journal")
    if journal_sheet:
        for r in _sheet_rows(wb, journal_sheet, ("id", "code", "name")):
            if r.get("code"):
                sheet_journal_codes.add(str(r["code"]).strip())
            if r.get("name"):
                sheet_journal_names.add(str(r["name"]).strip())

    standard_journal_codes = {"INV", "BILL", "MISC", "EXCH", "CABA", "TAX"}
    leftover_journals = env["account.journal"].search([
        ("company_id", "=", company.id),
        ("code", "not in", list(standard_journal_codes | sheet_journal_codes)),
    ])
    for j in list(leftover_journals):
        if j.name in sheet_journal_names:
            continue
        # JANGAN PERNAH hapus jurnal bawaan Odoo (memiliki External ID selain modul import ini atau id <= 8)
        is_system_journal = env["ir.model.data"].search_count([
            ("model", "=", "account.journal"),
            ("res_id", "=", j.id),
            ("module", "not in", (MODULE, "__import__")),
        ])
        if is_system_journal or j.id <= 8:
            continue

        try:
            with env.cr.savepoint():
                imds = env["ir.model.data"].search([("model", "=", "account.journal"), ("res_id", "=", j.id)])
                imds.unlink()
                j.unlink()
        except Exception:
            j.active = False

    # 4. Bersihkan akun COA custom lama yang TIDAK ADA di spreadsheet baru
    sheet_account_codes = set()
    account_sheet = find_sheet(wb, "account.account")
    if account_sheet:
        for r in _sheet_rows(wb, account_sheet, ("code",)):
            if r.get("code"):
                sheet_account_codes.add(str(r["code"]).strip())

    leftover_accounts = env["account.account"].with_context(active_test=False).search([
        ("company_ids", "in", company.id),
        ("code", "not in", list(sheet_account_codes)),
    ])
    journal_account_ids = set(
        env["account.journal"].search([("company_id", "=", company.id)]).mapped("default_account_id.id")
    ) | set(
        env["account.journal"].search([("company_id", "=", company.id)]).mapped("suspense_account_id.id")
    )

    for acc in leftover_accounts:
        # JANGAN PERNAH hapus akun bawaan Odoo (memiliki External ID dari modul sistem, misal 'account', 'l10n_id', 'base')
        is_system_account = env["ir.model.data"].search_count([
            ("model", "=", "account.account"),
            ("res_id", "=", acc.id),
            ("module", "not in", (MODULE, "__import__")),
        ])
        if is_system_account:
            # Akun bawaan sistem Odoo: jangan dihapus
            continue

        # Jangan hapus akun jika dipakai di konfigurasi perusahaan
        if acc.id in (
            company.transfer_account_id.id if company.transfer_account_id else 0,
            company.account_journal_suspense_account_id.id if company.account_journal_suspense_account_id else 0,
        ) or acc.id in journal_account_ids:
            acc.active = False
            continue

        try:
            with env.cr.savepoint():
                imds = env["ir.model.data"].search([("model", "=", "account.account"), ("res_id", "=", acc.id)])
                imds.unlink()
                acc.unlink()
        except Exception:
            acc.active = False


def run_import(env, wb, history=None, wb_formulas=None):
    """Run the full master-data import against an already-open workbook.
    Shared entry point for import_bundled_data() (a client module's
    post_init_hook), sf.import.history, and the Settings > Import Data Master wizard.

    wb_formulas: opsional, workbook yang sama tapi dibuka dengan
    data_only=False - kalau diberikan, dipakai untuk mendeteksi sel formula
    yang cache-nya kosong (lihat _warn_uncomputed_formulas) sebelum import
    berjalan.
    """
    company = env.ref("base.main_company")
    logger = ImportLogger(history)

    if wb_formulas is not None:
        _warn_uncomputed_formulas(wb_formulas, wb, logger=logger)

    # Validasi: Jika sudah ada transaksi/journal entry operasional yang diinput user, tolak import
    user_moves = _check_user_journal_entries(env, company)
    if user_moves:
        sample_names = ", ".join([str(m.name or m.ref or f"Draft #{m.id}") for m in user_moves[:5]])
        msg = (
            f"Import Ditolak: Ditemukan {len(user_moves)} transaksi/journal entry operasional "
            f"yang telah diinput oleh user ({sample_names}).\n\n"
            f"Untuk menjaga integritas dan validitas data akuntansi, proses import data master "
            f"hanya dapat dijalankan jika belum ada transaksi operasional yang dibuat manual.\n"
            f"Harap batalkan atau hapus journal entries tersebut terlebih dahulu jika Anda ingin "
            f"mengulang import data master dan saldo awal."
        )
        if logger:
            logger.log("account.move", 0, "VALIDATION", "Transaksi Operasional User", "error", msg)
            logger.flush()
        raise UserError(msg)

    # Bersihkan sisa data lama dari import sebelumnya atau database hasil clone
    _cleanup_previous_data(env, wb, company)

    balance_date = _read_opening_balance_date(wb)
    # account_opening_date perlu tanggal *awal* tahun buku, bukan tanggal
    # cutover-nya - Odoo men-tanggal-kan opening move sehari sebelum awal
    # tahun buku, jadi arahnya kebalik dari balance_date.
    fiscal_year_start = balance_date + timedelta(days=1)
    if fiscal_year_start.day != 1:
        _logger.warning(
            'OPENING_BALANCE_DATE (%s) menghasilkan awal tahun buku %s, bukan tanggal 1 - '
            "cek lagi kalau ini bukan disengaja (mis. salah pilih tanggal).",
            balance_date, fiscal_year_start,
        )

    company.write({"account_opening_date": fiscal_year_start})
    if company.account_opening_move_id and company.account_opening_move_id.state == "posted":
        company.account_opening_move_id.button_draft()
    if company.account_opening_move_id and company.account_opening_move_id.state == "draft":
        company.account_opening_move_id.date = balance_date

    analytic_percentage_digits = _read_decimal_accuracy_percentage_analytic(wb)
    _apply_analytic_percentage_precision(env, analytic_percentage_digits, logger=logger)

    _import_company(env, wb, logger=logger)
    _import_res_partner(env, wb, logger=logger)
    _import_account_account(env, wb, logger=logger)
    _import_partner_default_accounts(env, wb, logger=logger)
    _import_account_journal(env, wb, logger=logger)
    _import_account_asset_model(env, wb, logger=logger)
    _import_account_asset(env, wb, balance_date, logger=logger)
    _import_kas_bank(env, wb, balance_date, logger=logger)
    _import_partner_opening_balance(env, wb, balance_date, logger=logger)
    _import_opening_move(env, wb, "vendor_bill", "supplier_rank", balance_date, logger=logger)
    _import_opening_move(env, wb, "customer_invoice", "customer_rank", balance_date, logger=logger)
    _reconcile_liquidity_transfer(env)
    _import_account_analytic_plan(env, wb, logger=logger)
    _import_account_analytic_account(env, wb, logger=logger)
    _import_account_report(env, wb, logger=logger)
    _import_account_report_line(env, wb, logger=logger)
    if logger:
        logger.flush()
